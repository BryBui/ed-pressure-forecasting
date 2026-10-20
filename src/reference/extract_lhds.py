"""Extract geographic NSW Local Health Districts from a BHI workbook."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook


SNAPSHOT_PATTERN = re.compile(
    r"^ed_performance_(?P<year>\d{4})Q(?P<quarter>[1-4])_"
    r"(?P<timestamp>\d{8}T\d{6}Z)\.xlsx$"
)
REQUIRED_COLUMNS = {
    "reporting_level",
    "entity_id",
    "entity_longname",
    "entity_shortname",
}


@dataclass(frozen=True)
class LhdRecord:
    bhi_lhd_id: str
    lhd_name: str
    lhd_short_name: str
    source_workbook: str


def find_latest_bhi_workbook(raw_root: Path) -> Path:
    """Find the latest reporting period and snapshot in the BHI raw layer."""
    candidates: list[tuple[tuple[int, int, str], Path]] = []
    dataset_root = Path(raw_root) / "bhi" / "ed_performance"
    for path in dataset_root.glob("year=*/quarter=Q*/*.xlsx"):
        if match := SNAPSHOT_PATTERN.fullmatch(path.name):
            key = (
                int(match.group("year")),
                int(match.group("quarter")),
                match.group("timestamp"),
            )
            candidates.append((key, path))
    if not candidates:
        raise FileNotFoundError(f"No BHI ED snapshots found under {dataset_root}")
    return max(candidates, key=lambda item: item[0])[1]


def extract_geographic_lhds(workbook_path: Path) -> list[LhdRecord]:
    """Extract BHI LHD rows and exclude non-geographic specialty networks."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Overall results" not in workbook.sheetnames:
            raise ValueError("BHI workbook is missing the 'Overall results' sheet")
        rows = workbook["Overall results"].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("BHI 'Overall results' sheet is empty")
        indexes = {str(value).strip(): index for index, value in enumerate(headers)}
        missing = REQUIRED_COLUMNS - indexes.keys()
        if missing:
            raise ValueError(f"BHI workbook is missing columns: {sorted(missing)}")

        discovered: dict[str, LhdRecord] = {}
        for row in rows:
            if row[indexes["reporting_level"]] != "LHD":
                continue
            entity_id = str(row[indexes["entity_id"]] or "").strip()
            # BHI groups specialty networks at the LHD reporting level. NSW's
            # geographic LHD identifiers consistently end in "LHD".
            if not entity_id.endswith("LHD"):
                continue
            record = LhdRecord(
                bhi_lhd_id=entity_id,
                lhd_name=str(row[indexes["entity_longname"]] or "").strip(),
                lhd_short_name=str(
                    row[indexes["entity_shortname"]] or ""
                ).strip(),
                source_workbook=workbook_path.name,
            )
            previous = discovered.setdefault(entity_id, record)
            if previous != record:
                raise ValueError(f"Conflicting BHI definitions for {entity_id}")
    finally:
        workbook.close()

    if not discovered:
        raise ValueError("No geographic LHD records were found in the BHI workbook")
    return sorted(discovered.values(), key=lambda record: record.bhi_lhd_id)

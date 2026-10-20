"""Extract BHI hospital identities and their reported parent organisations."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook

from .extract_lhds import LhdRecord


PERIOD_PATTERN = re.compile(
    r"^(?P<start>Jan|Apr|Jul|Oct)-(?P<end>Mar|Jun|Sep|Dec) (?P<year>\d{4})$"
)
QUARTER_BY_START_MONTH = {"Jan": 1, "Apr": 2, "Jul": 3, "Oct": 4}
REQUIRED_COLUMNS = {
    "reporting_period",
    "reporting_level",
    "entity_id",
    "entity_longname",
    "entity_shortname",
    "lhd_shortname",
}


@dataclass(frozen=True)
class HospitalRecord:
    bhi_hospital_id: str
    hospital_name: str
    hospital_short_name: str
    bhi_parent_name: str
    bhi_parent_type: str
    geographic_lhd_id: str
    geographic_lhd_name: str
    first_reported_quarter: str
    latest_reported_quarter: str
    present_in_latest_quarter: str
    source_workbook: str


def _period_key(value: str) -> tuple[int, int]:
    match = PERIOD_PATTERN.fullmatch(value)
    if not match:
        raise ValueError(f"Unrecognised BHI reporting period: {value!r}")
    return int(match.group("year")), QUARTER_BY_START_MONTH[match.group("start")]


def _quarter_text(value: str) -> str:
    year, quarter = _period_key(value)
    return f"{year}Q{quarter}"


def extract_hospitals(
    workbook_path: Path,
    geographic_lhds: list[LhdRecord],
) -> list[HospitalRecord]:
    """Extract every hospital represented in BHI, including historical ones."""
    lhd_by_short_name = {record.lhd_short_name: record for record in geographic_lhds}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Overall results" not in workbook.sheetnames:
            raise ValueError("BHI workbook is missing the 'Overall results' sheet")
        rows = workbook["Overall results"].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("BHI 'Overall results' sheet is empty")
        indexes = {str(value).strip(): index for index, value in enumerate(headers)}
        missing = REQUIRED_COLUMNS - indexes.keys()
        if missing:
            raise ValueError(f"BHI workbook is missing columns: {sorted(missing)}")

        identities: dict[str, tuple[str, str, str]] = {}
        periods: dict[str, set[str]] = {}
        all_hospital_periods: set[str] = set()
        for row in rows:
            if row[indexes["reporting_level"]] != "Hospital":
                continue
            hospital_id = str(row[indexes["entity_id"]] or "").strip()
            identity = (
                str(row[indexes["entity_longname"]] or "").strip(),
                str(row[indexes["entity_shortname"]] or "").strip(),
                str(row[indexes["lhd_shortname"]] or "").strip(),
            )
            if not hospital_id or not all(identity):
                raise ValueError("BHI contains a hospital with incomplete identity fields")
            if hospital_id in identities and identities[hospital_id] != identity:
                raise ValueError(f"Conflicting BHI definitions for {hospital_id}")
            identities[hospital_id] = identity
            reporting_period = str(row[indexes["reporting_period"]] or "").strip()
            _period_key(reporting_period)
            periods.setdefault(hospital_id, set()).add(reporting_period)
            all_hospital_periods.add(reporting_period)
    finally:
        workbook.close()

    if not identities:
        raise ValueError("No hospital records were found in the BHI workbook")
    latest_period = max(all_hospital_periods, key=_period_key)

    records: list[HospitalRecord] = []
    for hospital_id, (name, short_name, parent_name) in identities.items():
        reported_periods = periods[hospital_id]
        first = min(reported_periods, key=_period_key)
        latest = max(reported_periods, key=_period_key)
        geographic_lhd = lhd_by_short_name.get(parent_name)
        records.append(
            HospitalRecord(
                bhi_hospital_id=hospital_id,
                hospital_name=name,
                hospital_short_name=short_name,
                bhi_parent_name=parent_name,
                bhi_parent_type=(
                    "geographic_lhd" if geographic_lhd else "specialty_network"
                ),
                geographic_lhd_id=(
                    geographic_lhd.bhi_lhd_id if geographic_lhd else ""
                ),
                geographic_lhd_name=(
                    geographic_lhd.lhd_name if geographic_lhd else ""
                ),
                first_reported_quarter=_quarter_text(first),
                latest_reported_quarter=_quarter_text(latest),
                present_in_latest_quarter=str(latest_period in reported_periods).lower(),
                source_workbook=workbook_path.name,
            )
        )
    return sorted(records, key=lambda record: record.bhi_hospital_id)

"""Discover, validate, and ingest ABS SA1 SEIFA index workbooks."""

from __future__ import annotations

from collections.abc import Callable, Iterator
from dataclasses import asdict, dataclass
from datetime import datetime
from hashlib import sha256
from html.parser import HTMLParser
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from typing import Any
from urllib.parse import unquote, urljoin, urlsplit

from openpyxl import load_workbook

from .storage import build_snapshot_path


LATEST_RELEASE_PAGE = (
    "https://www.abs.gov.au/statistics/people/people-and-communities/"
    "socio-economic-indexes-areas-seifa-australia/latest-release"
)
ABS_HOSTS = {"abs.gov.au", "www.abs.gov.au"}
WORKBOOK_PATTERN = re.compile(
    r"^Statistical Area Level 1, Indexes, SEIFA (?P<year>\d{4})\.xlsx$",
    re.IGNORECASE,
)
XLSX_CONTENT_TYPES = {
    "application/octet-stream",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
REQUIRED_SHEETS = {"Contents", *(f"Table {number}" for number in range(1, 8))}


class ReleaseLinkParser(HTMLParser):
    """Collect links from an ABS SEIFA release page."""

    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag != "a":
            return
        if href := dict(attrs).get("href"):
            self.links.append(href)


class ReleaseStatusParser(ReleaseLinkParser):
    """Collect the release marker and reference period from the ABS page."""

    def __init__(self) -> None:
        super().__init__()
        self.reference_period_parts: list[str] = []
        self.latest_release_parts: list[str] = []
        self._reference_depth = 0
        self._reference_item_depth = 0
        self._capture_latest_release = False

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        super().handle_starttag(tag, attrs)
        classes = set((dict(attrs).get("class") or "").split())

        if tag == "div":
            if self._reference_depth:
                self._reference_depth += 1
            elif "field--name-field-abs-reference-period" in classes:
                self._reference_depth = 1
            if self._reference_depth and "field__item" in classes:
                self._reference_item_depth = self._reference_depth

        if tag == "span" and "latest-release__text" in classes:
            self._capture_latest_release = True

    def handle_endtag(self, tag: str) -> None:
        if tag == "span":
            self._capture_latest_release = False
        if tag == "div" and self._reference_depth:
            if self._reference_depth == self._reference_item_depth:
                self._reference_item_depth = 0
            self._reference_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._reference_item_depth:
            self.reference_period_parts.append(data)
        if self._capture_latest_release:
            self.latest_release_parts.append(data)


@dataclass(frozen=True)
class SeifaWorkbook:
    """An official ABS SA1 index workbook advertised for a Census year."""

    url: str
    source_filename: str
    census_year: int

    @property
    def reporting_period(self) -> str:
        return str(self.census_year)


@dataclass(frozen=True)
class SeifaReleasePage:
    """Release signals and SA1 download advertised on the ABS page."""

    reference_period: int
    latest_release_text: str | None
    workbook: SeifaWorkbook


@dataclass(frozen=True)
class IngestionResult:
    """Serializable SEIFA ingestion outcome returned to orchestration."""

    status: str
    path: str
    source_url: str
    sha256: str
    reporting_period: str
    geography: str = "SA1"

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


def _require_official_abs_url(url: str, label: str) -> None:
    parsed = urlsplit(url)
    if parsed.scheme != "https" or parsed.hostname not in ABS_HOSTS:
        raise ValueError(f"{label} must be an HTTPS URL on abs.gov.au: {url!r}")


def workbook_from_url(url: str) -> SeifaWorkbook | None:
    """Parse an official SA1 index workbook URL."""
    _require_official_abs_url(url, "workbook URL")
    filename = Path(unquote(urlsplit(url).path)).name
    match = WORKBOOK_PATTERN.fullmatch(filename)
    if not match:
        return None
    return SeifaWorkbook(
        url=url,
        source_filename=filename,
        census_year=int(match.group("year")),
    )


def discover_sa1_workbook(
    html: str, page_url: str = LATEST_RELEASE_PAGE
) -> SeifaWorkbook:
    """Find the unique SA1 index workbook advertised by an ABS release page."""
    _require_official_abs_url(page_url, "release page URL")
    parser = ReleaseLinkParser()
    parser.feed(html)

    candidates: dict[str, SeifaWorkbook] = {}
    for link in parser.links:
        url = urljoin(page_url, link)
        try:
            workbook = workbook_from_url(url)
        except ValueError:
            continue
        if workbook is not None:
            candidates[url] = workbook

    if len(candidates) != 1:
        raise RuntimeError(
            "Could not find a unique ABS Statistical Area Level 1 SEIFA "
            "indexes workbook link"
        )
    return next(iter(candidates.values()))


def inspect_latest_release_page(
    html: str, page_url: str = LATEST_RELEASE_PAGE
) -> SeifaReleasePage:
    """Read the release signals and SA1 workbook from the official ABS page."""
    _require_official_abs_url(page_url, "release page URL")
    parser = ReleaseStatusParser()
    parser.feed(html)

    reference_period_text = " ".join(parser.reference_period_parts).strip()
    if not re.fullmatch(r"\d{4}", reference_period_text):
        raise RuntimeError(
            "Could not find a four-digit SEIFA reference period on the ABS page"
        )

    workbook = discover_sa1_workbook(html, page_url)
    reference_period = int(reference_period_text)
    if workbook.census_year != reference_period:
        raise RuntimeError(
            "ABS SEIFA reference period does not match the advertised SA1 "
            f"workbook year: {reference_period} != {workbook.census_year}"
        )

    latest_release_text = " ".join(parser.latest_release_parts).strip() or None
    return SeifaReleasePage(
        reference_period=reference_period,
        latest_release_text=latest_release_text,
        workbook=workbook,
    )


def _default_get(url: str, **kwargs: Any) -> Any:
    try:
        from curl_cffi import requests
    except ImportError as error:
        raise RuntimeError(
            "curl-cffi is required for live ABS downloads; install the project "
            "dependencies first"
        ) from error
    return requests.get(url, impersonate="chrome", **kwargs)


def resolve_sa1_workbook(
    release_url: str = LATEST_RELEASE_PAGE,
    *,
    get: Callable[..., Any] = _default_get,
) -> SeifaWorkbook:
    """Resolve either an ABS workbook URL or an ABS release-page URL."""
    _require_official_abs_url(release_url, "release URL")
    if workbook := workbook_from_url(release_url):
        return workbook

    response = get(release_url, timeout=60)
    try:
        response.raise_for_status()
        return discover_sa1_workbook(response.text, release_url)
    finally:
        response.close()


def _normalised_cells(values: tuple[Any, ...]) -> tuple[str, ...]:
    return tuple("" if value is None else str(value).strip() for value in values)


def _require_header(
    worksheet: Any, row_number: int, expected_prefix: tuple[str, ...]
) -> None:
    row = next(
        worksheet.iter_rows(
            min_row=row_number,
            max_row=row_number,
            max_col=len(expected_prefix),
            values_only=True,
        )
    )
    actual = _normalised_cells(row)
    if actual != expected_prefix:
        raise RuntimeError(
            f"Unexpected {worksheet.title} row {row_number} header: {actual!r}"
        )


def _require_sa1_data(worksheet: Any, row_number: int) -> None:
    value = worksheet.cell(row=row_number, column=1).value
    code = str(value).strip() if value is not None else ""
    if not (len(code) == 11 and code.isdigit()):
        raise RuntimeError(
            f"{worksheet.title} does not contain an 11-digit SA1 code at "
            f"row {row_number}"
        )


def validate_sa1_workbook(path: Path, census_year: int) -> None:
    """Validate the geography, year, sheets, headers, and first data records."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise RuntimeError("ABS response is not a readable XLSX workbook") from error

    try:
        missing = REQUIRED_SHEETS.difference(workbook.sheetnames)
        if missing:
            raise RuntimeError(
                "ABS SA1 workbook is missing required sheets: "
                + ", ".join(sorted(missing))
            )

        title = str(workbook["Contents"].cell(row=2, column=1).value or "")
        if "SEIFA" not in title or str(census_year) not in title:
            raise RuntimeError(
                "ABS workbook title does not match the advertised Census year"
            )

        _require_header(
            workbook["Table 1"],
            6,
            (
                f"{census_year} Statistical Area Level 1 (SA1)",
                "Score",
                "Decile",
                "Score",
                "Decile",
                "Score",
                "Decile",
                "Score",
                "Decile",
                "Usual Resident Population",
            ),
        )
        index_header = (
            f"{census_year} Statistical Area Level 1 (SA1)",
            "Usual Resident Population",
            "Score",
        )
        for sheet_name in ("Table 2", "Table 3", "Table 4", "Table 5"):
            _require_header(workbook[sheet_name], 6, index_header)
        _require_header(
            workbook["Table 6"],
            6,
            (
                f"{census_year} Statistical Area Level 1 (SA1)",
                "Usual Resident Population",
                "IRSD",
                "IRSAD",
                "IER",
                "IEO",
            ),
        )
        _require_header(
            workbook["Table 7"],
            5,
            (f"{census_year} Statistical Area Level 1 (SA1) Code",),
        )
        for sheet_name in ("Table 1", "Table 2", "Table 3", "Table 4", "Table 5"):
            _require_sa1_data(workbook[sheet_name], 7)
        _require_sa1_data(workbook["Table 7"], 6)
    finally:
        workbook.close()


def _response_chunks(response: Any) -> Iterator[bytes]:
    for chunk in response.iter_content(chunk_size=1024 * 1024):
        if chunk:
            yield chunk


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _matching_snapshot(directory: Path, checksum: str) -> Path | None:
    for candidate in sorted(directory.glob("seifa_*.xlsx")):
        if candidate.name.endswith(".part.xlsx"):
            continue
        if _sha256_file(candidate) == checksum:
            return candidate
    return None


def _download_to_temporary_file(
    workbook: SeifaWorkbook,
    directory: Path,
    get: Callable[..., Any],
) -> tuple[Path, str]:
    directory.mkdir(parents=True, exist_ok=True)
    temporary = NamedTemporaryFile(
        mode="wb", suffix=".part.xlsx", prefix="seifa_", dir=directory, delete=False
    )
    temporary_path = Path(temporary.name)
    digest = sha256()
    response = None
    try:
        response = get(workbook.url, stream=True, timeout=180)
        response.raise_for_status()
        content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
        if content_type not in XLSX_CONTENT_TYPES:
            raise RuntimeError(
                "Expected an ABS Excel workbook, but received "
                f"{content_type or 'no content type'}"
            )
        for chunk in _response_chunks(response):
            temporary.write(chunk)
            digest.update(chunk)
        temporary.close()
        validate_sa1_workbook(temporary_path, workbook.census_year)
        return temporary_path, digest.hexdigest()
    except Exception:
        temporary.close()
        temporary_path.unlink(missing_ok=True)
        raise
    finally:
        if response is not None:
            response.close()


def ingest_sa1_release(
    raw_root: Path,
    ingested_at: datetime,
    *,
    release_url: str = LATEST_RELEASE_PAGE,
    get: Callable[..., Any] = _default_get,
) -> IngestionResult:
    """Ingest an advertised SA1 workbook once, preserving later revisions."""
    workbook = resolve_sa1_workbook(release_url, get=get)
    destination = build_snapshot_path(
        raw_root,
        source="abs",
        dataset="seifa",
        partitions={"census_year": str(workbook.census_year)},
        reporting_period=workbook.reporting_period,
        ingested_at=ingested_at,
        extension="xlsx",
    )
    temporary_path, checksum = _download_to_temporary_file(
        workbook, destination.parent, get
    )
    try:
        if existing := _matching_snapshot(destination.parent, checksum):
            return IngestionResult(
                status="unchanged",
                path=str(existing),
                source_url=workbook.url,
                sha256=checksum,
                reporting_period=workbook.reporting_period,
            )

        try:
            os.link(temporary_path, destination)
        except FileExistsError:
            if _sha256_file(destination) != checksum:
                raise RuntimeError(
                    f"Snapshot path already exists with different content: {destination}"
                )
            status = "unchanged"
        else:
            status = "created"

        return IngestionResult(
            status=status,
            path=str(destination),
            source_url=workbook.url,
            sha256=checksum,
            reporting_period=workbook.reporting_period,
        )
    finally:
        temporary_path.unlink(missing_ok=True)


def ingest_latest_sa1_if_new_period(
    raw_root: Path,
    ingested_at: datetime,
    *,
    get: Callable[..., Any] = _default_get,
) -> IngestionResult:
    """Poll the ABS release page and ingest only a newly advertised period."""
    response = get(LATEST_RELEASE_PAGE, timeout=60)
    try:
        response.raise_for_status()
        release = inspect_latest_release_page(response.text)
    finally:
        response.close()

    if release.latest_release_text != "Latest release":
        raise RuntimeError(
            "ABS SEIFA page contract changed: expected the 'Latest release' "
            f"marker, received {release.latest_release_text!r}"
        )

    partition = (
        raw_root / "abs" / "seifa" / f"census_year={release.reference_period}"
    )
    existing_snapshots = sorted(
        path
        for path in partition.glob(f"seifa_{release.reference_period}_*.xlsx")
        if not path.name.endswith(".part.xlsx")
    )
    if existing_snapshots:
        existing = existing_snapshots[-1]
        return IngestionResult(
            status="unchanged",
            path=str(existing),
            source_url=release.workbook.url,
            sha256=_sha256_file(existing),
            reporting_period=str(release.reference_period),
        )

    return ingest_sa1_release(
        raw_root,
        ingested_at,
        release_url=release.workbook.url,
        get=get,
    )

"""Export the latest local BHI ED workbook to reference and silver CSVs."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
import re
from shutil import copyfileobj
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

DEFAULT_RAW = Path("data/raw/bhi/ed_performance")
DEFAULT_REFERENCE = Path("data/reference")
DEFAULT_SILVER = Path("data/silver/bhi")
SNAPSHOT_PATTERN = re.compile(r"ed_performance_(\d{4}Q[1-4])_(\d{8}T\d{6}Z)\.xlsx")


def latest_workbook(raw_dir: Path) -> Path:
    """Select by reporting quarter, then ingestion timestamp; ignore partial files."""
    candidates = [
        (match.groups(), str(path), path)
        for path in raw_dir.rglob("*.xlsx")
        if path.is_file() and (match := SNAPSHOT_PATTERN.fullmatch(path.name))
    ]
    if not candidates:
        raise FileNotFoundError(f"No BHI ED workbook snapshots found under {raw_dir}")
    return max(candidates)[2]


def export_workbook(
    workbook_path: Path,
    reference_dir: Path = DEFAULT_REFERENCE,
    silver_dir: Path = DEFAULT_SILVER,
) -> dict[Path, int]:
    """Export cached cell values as UTF-8 CSVs, preserving results sheet rows.

    Fields contains a presentation preamble: export only its Field/Name/Description
    table. Stage all sheets before publishing so a malformed sheet leaves existing
    outputs intact. Returned counts exclude each CSV's header.
    """
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Fields" not in workbook.sheetnames:
            raise ValueError("BHI workbook is missing the Fields sheet")
        exports = {}
        for name in workbook.sheetnames:
            if name == "Metadata":
                continue
            filename = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
            if not filename:
                raise ValueError(f"Cannot create a CSV filename for sheet {name!r}")
            destination = (
                reference_dir / "bhi_field_definitions.csv"
                if name == "Fields" else silver_dir / f"{filename}.csv"
            )
            if destination.resolve() == workbook_path.resolve():
                raise ValueError("Output must differ from the input workbook")
            if destination.resolve() in {p.resolve() for p in exports.values()}:
                raise ValueError(f"Duplicate CSV destination for sheet {name!r}")
            exports[name] = destination

        counts = {}
        with TemporaryDirectory(prefix="bhi_ed_export_") as staging:
            staged = []
            for index, (name, destination) in enumerate(exports.items()):
                rows = iter(workbook[name].values)
                if name == "Fields":
                    for row in rows:
                        for column in range(len(row) - 2):
                            if row[column:column + 3] == ("Field", "Name", "Description"):
                                break
                        else:
                            continue
                        break
                    else:
                        raise ValueError("Fields sheet has no Field/Name/Description header")
                    rows = (row[column:column + 3] for row in rows if any(row[column:column + 3]))
                    header = ("Field", "Name", "Description")
                else:
                    header = next(rows, None)
                    if header is None:
                        raise ValueError(f"Empty results sheet: {name}")
                temporary = Path(staging) / f"{index}.csv"
                count = 0
                with temporary.open("w", newline="", encoding="utf-8") as output:
                    writer = csv.writer(output)
                    writer.writerow(header)
                    for row in rows:
                        writer.writerow(row)
                        count += 1
                staged.append((temporary, destination))
                counts[destination] = count
            for temporary, destination in staged:
                destination.parent.mkdir(parents=True, exist_ok=True)
                # Staging may be on a different filesystem from the project.
                with temporary.open("rb") as source, destination.open("wb") as output:
                    copyfileobj(source, output)
        return counts
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW)
    parser.add_argument("--reference-dir", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--silver-dir", type=Path, default=DEFAULT_SILVER)
    args = parser.parse_args()
    try:
        source = latest_workbook(args.raw_dir)
        counts = export_workbook(source, args.reference_dir, args.silver_dir)
    except (FileNotFoundError, ValueError) as error:
        parser.error(str(error))
    print(f"Source: {source}")
    for path, count in counts.items():
        print(f"Wrote {count} rows to {path}")


if __name__ == "__main__":
    main()

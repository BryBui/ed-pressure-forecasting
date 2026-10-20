"""Aggregate 2021 SA1 SEIFA scores into descriptive LHD summaries.

Only SA1s with all four scores and a population are included in any calculation.
Means describe constituent SA1 scores; they are not official ABS LHD indexes.
Population is allocated wholly to the LHD selected by the correspondence.
"""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path

from openpyxl import load_workbook

INDEXES = ("irsd", "irsad", "ier", "ieo")
DEFAULT_MAPPING = Path("data/reference/sa1_to_lhd.csv")
DEFAULT_OUTPUT = Path("data/silver/linking/lhd_socioeconomic.csv")


def _number(value: object) -> float | None:
    if value is None or str(value).strip() in {"", "-", "..", "n.a."}:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid SEIFA numeric value: {value!r}") from error
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"Invalid SEIFA numeric value: {value!r}")
    return number


def read_scores(path: Path) -> dict[str, tuple[float | None, ...]]:
    """Read Table 1 once, retaining missing scores for explicit exclusion."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Table 1"]
        header = next(sheet.iter_rows(min_row=6, max_row=6, max_col=10, values_only=True))
        if header != (
            "2021 Statistical Area Level 1 (SA1)", "Score", "Decile",
            "Score", "Decile", "Score", "Decile", "Score", "Decile",
            "Usual Resident Population",
        ):
            raise ValueError("Expected 2021 SEIFA Table 1 headers")
        scores = {}
        for row in sheet.iter_rows(min_row=7, max_col=10, values_only=True):
            code = str(row[0]).strip()
            if not (len(code) == 11 and code.isdigit()):
                continue  # Blank rows and ABS footnotes are not SA1 records.
            if code in scores:
                raise ValueError(f"Duplicate SEIFA SA1: {code}")
            scores[code] = tuple(_number(row[i]) for i in (9, 1, 3, 5, 7))
        if not scores:
            raise ValueError("SEIFA Table 1 contains no SA1 records")
        return scores
    finally:
        workbook.close()


def aggregate_lhds(mapping_path: Path, seifa_path: Path) -> list[dict]:
    """Return one row per mapped LHD, including LHDs with no usable scores."""
    return _aggregate_groups(mapping_path, seifa_path, "lhd_code", "lhd_name")


def _aggregate_groups(
    mapping_path: Path, seifa_path: Path, id_field: str, name_field: str,
) -> list[dict]:
    """Shared complete-case SEIFA aggregation for geographic correspondences."""
    scores = read_scores(seifa_path)
    groups = {}
    seen = set()
    with mapping_path.open(newline="", encoding="utf-8-sig") as source:
        reader = csv.DictReader(source)
        if not {"sa1_code_2021", id_field, name_field}.issubset(reader.fieldnames or []):
            raise ValueError(f"Mapping requires sa1_code_2021, {id_field}, {name_field}")
        for record in reader:
            code, lhd, name = (record[k].strip() for k in ("sa1_code_2021", id_field, name_field))
            if len(code) != 11 or not code.isdigit() or not lhd or not name:
                raise ValueError(f"Invalid mapping record: {record}")
            if code in seen:
                raise ValueError(f"Duplicate mapped SA1: {code}")
            seen.add(code)
            group = groups.setdefault(lhd, {"name": name, "mapped": 0, "values": []})
            if group["name"] != name:
                raise ValueError(f"Conflicting names for group {lhd}")
            group["mapped"] += 1
            values = scores.get(code)
            if values is not None and all(v is not None for v in values):
                group["values"].append(values)
    if not groups:
        raise ValueError("Mapping contains no SA1 records")
    result = []
    for lhd, group in sorted(groups.items()):
        values = group["values"]
        population = math.fsum(v[0] for v in values)
        row = dict(census_year=2021,
                   mapped_sa1_count=group["mapped"], included_sa1_count=len(values),
                   excluded_sa1_count=group["mapped"] - len(values),
                   population_sum=population)
        row = {id_field: lhd, name_field: group["name"], **row}
        for position, index in enumerate(INDEXES, 1):
            row[f"{index}_mean"] = math.fsum(v[position] for v in values) / len(values) if values else None
            row[f"{index}_population_weighted_mean"] = (
                math.fsum(v[position] * v[0] for v in values) / population
                if population else None
            )
        result.append(row)
    return result


def build_lhd_socioeconomic(mapping_path: Path, seifa_path: Path, output_path: Path) -> int:
    """Validate inputs and atomically write the LHD summary CSV."""
    if output_path.resolve() in {mapping_path.resolve(), seifa_path.resolve()}:
        raise ValueError("Output must differ from input files")
    rows = aggregate_lhds(mapping_path, seifa_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.tmp")
    try:
        with temporary.open("w", newline="", encoding="utf-8") as destination:
            writer = csv.DictWriter(destination, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        temporary.replace(output_path)
    finally:
        temporary.unlink(missing_ok=True)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    parser.add_argument("--seifa", type=Path, help="2021 ABS SEIFA workbook; defaults to latest local snapshot")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.seifa is None:
        snapshots = sorted(Path("data/raw/abs/seifa/census_year=2021").glob("seifa_2021_*.xlsx"))
        snapshots = [p for p in snapshots if not p.name.endswith(".part.xlsx")]
        if not snapshots:
            parser.error("No local 2021 SEIFA workbook found; supply --seifa")
        args.seifa = snapshots[-1]
    count = build_lhd_socioeconomic(args.mapping, args.seifa, args.output)
    print(f"Wrote {count} LHD summaries to {args.output} using {args.seifa}")


if __name__ == "__main__":
    main()

import csv
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from transform.bhi.bhi_ed_to_csv import export_workbook, latest_workbook


class BhiEdExportTests(unittest.TestCase):
    def test_latest_uses_period_then_ingestion_time(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            names = [
                "year=2025/ed_performance_2025Q4_20261020T120000Z.xlsx",
                "year=2026/ed_performance_2026Q2_20261018T113620Z.xlsx",
                "year=2026/ed_performance_2026Q2_20261019T113620Z.xlsx",
                "year=2026/ed_performance_2026Q3_20261020T113620Z.part.xlsx",
                "~$ed_performance_2026Q3_20261020T113620Z.xlsx",
            ]
            for name in names:
                path = root / name
                path.parent.mkdir(parents=True, exist_ok=True)
                path.touch()
            self.assertEqual(latest_workbook(root), root / names[2])

    def test_missing_workbook(self):
        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(FileNotFoundError, "No BHI ED workbook"):
                latest_workbook(Path(directory))

    def test_routes_sheets_and_preserves_values(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            workbook.active.title = "Metadata"
            workbook.active.append(["Do not export"])
            fields = workbook.create_sheet("Fields")
            fields.append([None, "Healthcare Quarterly"])
            fields.append([None, "Field", "Name", "Description", None])
            fields.append([None, "Notes", "note_text", 'Unicode – and commas, "quotes"'])
            for name in ("Overall results", "Results by triage"):
                sheet = workbook.create_sheet(name)
                sheet.append(["entity_id", "measure_value", "note_text"])
                sheet.append(["00123", 12.5, "line one\nline two"])
                sheet.append(["NSW", None, "–"])
            source = root / "source.xlsx"
            workbook.save(source)
            workbook.close()
            counts = export_workbook(source, root / "reference", root / "silver")
            self.assertEqual(sorted(p.name for p in counts), [
                "bhi_field_definitions.csv", "overall_results.csv", "results_by_triage.csv",
            ])
            with (root / "reference/bhi_field_definitions.csv").open(newline="", encoding="utf-8") as stream:
                self.assertEqual(list(csv.reader(stream)), [
                    ["Field", "Name", "Description"],
                    ["Notes", "note_text", 'Unicode – and commas, "quotes"'],
                ])
            for filename in ("overall_results.csv", "results_by_triage.csv"):
                with (root / "silver" / filename).open(newline="", encoding="utf-8") as stream:
                    self.assertEqual(list(csv.reader(stream)), [
                        ["entity_id", "measure_value", "note_text"],
                        ["00123", "12.5", "line one\nline two"],
                        ["NSW", "", "–"],
                    ])
            self.assertEqual(counts[root / "silver/overall_results.csv"], 2)
            self.assertFalse(list(root.rglob("metadata.csv")))

    def test_malformed_fields_does_not_replace_existing_csv(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = Workbook()
            workbook.active.title = "Overall results"
            workbook.active.append(["header"])
            workbook.create_sheet("Fields").append(["unexpected header"])
            source = root / "source.xlsx"
            workbook.save(source)
            workbook.close()
            silver = root / "silver"
            silver.mkdir()
            existing = silver / "overall_results.csv"
            existing.write_text("existing data\n")
            with self.assertRaisesRegex(ValueError, "no Field/Name/Description header"):
                export_workbook(source, root / "reference", silver)
            self.assertEqual(existing.read_text(), "existing data\n")


if __name__ == "__main__":
    unittest.main()

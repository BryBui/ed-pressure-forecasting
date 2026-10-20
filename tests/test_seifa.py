from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from ingestion.seifa import (
    LATEST_RELEASE_PAGE,
    discover_sa1_workbook,
    ingest_latest_sa1_if_new_period,
    ingest_sa1_release,
    inspect_latest_release_page,
    validate_sa1_workbook,
    workbook_from_url,
)


RELEASE_PAGE = (
    "https://www.abs.gov.au/statistics/people/people-and-communities/"
    "socio-economic-indexes-areas-seifa-australia/2021"
)
SA1_URL = RELEASE_PAGE + (
    "/Statistical%20Area%20Level%201%2C%20Indexes%2C%20SEIFA%202021.xlsx"
)
SA2_URL = RELEASE_PAGE + (
    "/Statistical%20Area%20Level%202%2C%20Indexes%2C%20SEIFA%202021.xlsx"
)
HTML = f'<a href="{SA2_URL}">SA2</a><a href="{SA1_URL}">SA1</a>'
LATEST_PAGE_HTML = f"""
<div class="abs-layout-published-date field
 field--name-field-abs-reference-period field--type-abs-reference-field-type
 field--label-inline">
  <div class="field__label">Reference period</div>
  <div class="field__item">2021</div>
</div>
<div class="field field--name-dynamic-block-fieldnode-latest-release">
  <div class="abs_latest_release_wrapper">
    <span class="latest-release__text">Latest release</span>
    <a href="#data-downloads">Data download</a>
  </div>
</div>
{HTML}
"""


def workbook_bytes(*, valid_sa1: bool = True, marker: str = "original") -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    contents = workbook.create_sheet("Contents")
    contents.cell(row=2, column=1, value="Socio-Economic Indexes for Australia (SEIFA), 2021")

    table_1 = workbook.create_sheet("Table 1")
    table_1.append([])
    table_1.cell(row=4, column=1, value=f"Table 1 SA1 SEIFA Summary 2021 {marker}")
    summary_header = [
        "2021 Statistical Area Level 1 (SA1)",
        "Score",
        "Decile",
        "Score",
        "Decile",
        "Score",
        "Decile",
        "Score",
        "Decile",
        "Usual Resident Population",
    ]
    for column, value in enumerate(summary_header, 1):
        table_1.cell(row=6, column=column, value=value)
    table_1.cell(row=7, column=1, value=10102100701 if valid_sa1 else 101021007)

    index_header = [
        "2021 Statistical Area Level 1 (SA1)",
        "Usual Resident Population",
        "Score",
    ]
    for number in range(2, 6):
        sheet = workbook.create_sheet(f"Table {number}")
        for column, value in enumerate(index_header, 1):
            sheet.cell(row=6, column=column, value=value)
        sheet.cell(row=7, column=1, value=10102100701)

    table_6 = workbook.create_sheet("Table 6")
    excluded_header = [
        "2021 Statistical Area Level 1 (SA1)",
        "Usual Resident Population",
        "IRSD",
        "IRSAD",
        "IER",
        "IEO",
    ]
    for column, value in enumerate(excluded_header, 1):
        table_6.cell(row=6, column=column, value=value)

    table_7 = workbook.create_sheet("Table 7")
    table_7.cell(
        row=5,
        column=1,
        value="2021 Statistical Area Level 1 (SA1) Code",
    )
    table_7.cell(row=6, column=1, value=10102100701)
    workbook.create_sheet("Explanatory Notes")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class FakeResponse:
    def __init__(
        self,
        *,
        text: str = "",
        body: bytes = b"",
        content_type: str = "text/html",
    ) -> None:
        self.text = text
        self.body = body
        self.headers = {"content-type": content_type}
        self.closed = False

    def raise_for_status(self) -> None:
        return None

    def iter_content(self, chunk_size: int):
        del chunk_size
        yield self.body

    def close(self) -> None:
        self.closed = True


class SeifaTests(unittest.TestCase):
    def test_reads_reference_period_and_latest_release_marker(self) -> None:
        release = inspect_latest_release_page(
            LATEST_PAGE_HTML, LATEST_RELEASE_PAGE
        )

        self.assertEqual(release.reference_period, 2021)
        self.assertEqual(release.latest_release_text, "Latest release")
        self.assertEqual(release.workbook.url, SA1_URL)

    def test_watcher_reuses_existing_period_without_downloading(self) -> None:
        calls: list[str] = []

        def fake_get(url: str, **kwargs):
            del kwargs
            calls.append(url)
            if url == LATEST_RELEASE_PAGE:
                return FakeResponse(text=LATEST_PAGE_HTML)
            raise AssertionError(f"Unexpected workbook download: {url}")

        with TemporaryDirectory() as directory:
            raw_root = Path(directory)
            existing = (
                raw_root
                / "abs/seifa/census_year=2021"
                / "seifa_2021_20261019T090000Z.xlsx"
            )
            existing.parent.mkdir(parents=True)
            existing.write_bytes(b"existing snapshot")

            result = ingest_latest_sa1_if_new_period(
                raw_root,
                datetime(2026, 10, 20, 9, tzinfo=timezone.utc),
                get=fake_get,
            )

        self.assertEqual(result.status, "unchanged")
        self.assertEqual(Path(result.path), existing)
        self.assertEqual(calls, [LATEST_RELEASE_PAGE])

    def test_watcher_rejects_a_missing_latest_release_marker(self) -> None:
        html = LATEST_PAGE_HTML.replace(
            '<span class="latest-release__text">Latest release</span>', ""
        )

        def fake_get(url: str, **kwargs):
            del url, kwargs
            return FakeResponse(text=html)

        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(RuntimeError, "Latest release"):
                ingest_latest_sa1_if_new_period(
                    Path(directory),
                    datetime(2026, 10, 20, 9, tzinfo=timezone.utc),
                    get=fake_get,
                )

    def test_watcher_downloads_a_new_reference_period(self) -> None:
        body = workbook_bytes()

        def fake_get(url: str, **kwargs):
            del kwargs
            if url == LATEST_RELEASE_PAGE:
                return FakeResponse(text=LATEST_PAGE_HTML)
            if url == SA1_URL:
                return FakeResponse(
                    body=body,
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            raise AssertionError(f"Unexpected URL: {url}")

        with TemporaryDirectory() as directory:
            result = ingest_latest_sa1_if_new_period(
                Path(directory),
                datetime(2026, 10, 20, 9, tzinfo=timezone.utc),
                get=fake_get,
            )

            self.assertEqual(result.status, "created")
            self.assertTrue(Path(result.path).is_file())

    def test_discovers_only_the_sa1_indexes_workbook(self) -> None:
        workbook = discover_sa1_workbook(HTML, RELEASE_PAGE)

        self.assertEqual(workbook.url, SA1_URL)
        self.assertEqual(workbook.census_year, 2021)

    def test_rejects_a_non_abs_workbook_host(self) -> None:
        with self.assertRaisesRegex(ValueError, "abs.gov.au"):
            workbook_from_url(
                "https://example.test/Statistical%20Area%20Level%201%2C%20"
                "Indexes%2C%20SEIFA%202021.xlsx"
            )

    def test_rejects_a_workbook_without_sa1_records(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.xlsx"
            path.write_bytes(workbook_bytes(valid_sa1=False))

            with self.assertRaisesRegex(RuntimeError, "11-digit SA1 code"):
                validate_sa1_workbook(path, 2021)

    def test_ingestion_is_partitioned_and_duplicate_safe(self) -> None:
        body = workbook_bytes()

        def fake_get(url: str, **kwargs):
            del kwargs
            if url == RELEASE_PAGE:
                return FakeResponse(text=HTML)
            if url == SA1_URL:
                return FakeResponse(
                    body=body,
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            raise AssertionError(f"Unexpected URL: {url}")

        with TemporaryDirectory() as directory:
            raw_root = Path(directory)
            first = ingest_sa1_release(
                raw_root,
                datetime(2026, 10, 19, 9, tzinfo=timezone.utc),
                release_url=RELEASE_PAGE,
                get=fake_get,
            )
            later = ingest_sa1_release(
                raw_root,
                datetime(2026, 10, 20, 9, tzinfo=timezone.utc),
                release_url=RELEASE_PAGE,
                get=fake_get,
            )

            expected = (
                raw_root
                / "abs/seifa/census_year=2021"
                / "seifa_2021_20261019T090000Z.xlsx"
            )
            self.assertEqual(first.status, "created")
            self.assertEqual(Path(first.path), expected)
            self.assertEqual(later.status, "unchanged")
            self.assertEqual(Path(later.path), expected)
            self.assertEqual(first.geography, "SA1")

    def test_revised_workbook_creates_another_snapshot(self) -> None:
        revision = {"body": workbook_bytes()}

        def fake_get(url: str, **kwargs):
            del kwargs
            if url == SA1_URL:
                return FakeResponse(
                    body=revision["body"],
                    content_type="application/octet-stream",
                )
            raise AssertionError(f"Unexpected URL: {url}")

        with TemporaryDirectory() as directory:
            raw_root = Path(directory)
            first = ingest_sa1_release(
                raw_root,
                datetime(2026, 10, 19, 9, tzinfo=timezone.utc),
                release_url=SA1_URL,
                get=fake_get,
            )
            revision["body"] = workbook_bytes(marker="revision")
            revised = ingest_sa1_release(
                raw_root,
                datetime(2026, 10, 20, 9, tzinfo=timezone.utc),
                release_url=SA1_URL,
                get=fake_get,
            )

            self.assertEqual(first.status, "created")
            self.assertEqual(revised.status, "created")
            self.assertNotEqual(first.path, revised.path)


if __name__ == "__main__":
    unittest.main()
